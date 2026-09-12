from pathlib import Path
from typing import Any, Dict, List, Union
import docx


def parse_docx(file_path: Union[str, Path]) -> Dict[str, Any]:
    """Parse Word documents (.docx) into a normalized document structure."""
    path = Path(file_path)
    doc = docx.Document(path)

    raw_paragraphs: List[str] = []

    # Extract text from standard document paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            raw_paragraphs.append(text)

    # Also extract text from tables if present
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                raw_paragraphs.append(row_text)

    paragraphs: List[Dict[str, Any]] = [
        {"index": idx, "text": para_text, "page": None}
        for idx, para_text in enumerate(raw_paragraphs)
    ]

    full_text = "\n\n".join(raw_paragraphs)

    return {
        "filename": path.name,
        "file_type": "docx",
        "text": full_text,
        "paragraphs": paragraphs,
        "metadata": {
            "page_count": None,
            "paragraph_count": len(paragraphs),
            "char_count": len(full_text),
        },
    }
